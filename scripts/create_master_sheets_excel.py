"""
=============================================================================
GBI BAITANI SUKABUMI — MASTER SPREADSHEET BUILDER
File: scripts/create_master_sheets_excel.py

Membangun file Excel "GBI_Baitani_Database_Master.xlsx" dengan 7 tab lengkap,
format header Warm Sanctuary (Emas #785600, teks putih tebal, auto-column width,
dan freeze pane), yang siap diunggah ke Google Drive dan langsung dikonversi
menjadi Google Sheets resmi gereja.
=============================================================================
"""

import os
import csv
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

BASE_DIR = os.path.dirname(os.path.dirname(__file__))
SEEDS_DIR = os.path.join(BASE_DIR, "data", "seeds")
OUTPUT_PROJECT = os.path.join(BASE_DIR, "data", "GBI_Baitani_Database_Master.xlsx")
OUTPUT_DOWNLOADS = r"C:\Users\jtoem\Downloads\GBI_Baitani_Database_Master.xlsx"

# Warm Sanctuary Styling Constants
HEADER_FILL = PatternFill(start_color="785600", end_color="785600", fill_type="solid")
HEADER_FONT = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
ROW_FONT = Font(name="Calibri", size=10, color="231F1A")
ZEBRA_FILL = PatternFill(start_color="FBF9F6", end_color="FBF9F6", fill_type="solid")
WHITE_FILL = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")

THIN_BORDER = Border(
    left=Side(style='thin', color='E8E1D5'),
    right=Side(style='thin', color='E8E1D5'),
    top=Side(style='thin', color='E8E1D5'),
    bottom=Side(style='thin', color='E8E1D5')
)

def read_csv(filename):
    path = os.path.join(SEEDS_DIR, filename)
    rows = []
    if os.path.exists(path):
        with open(path, 'r', encoding='utf-8') as f:
            reader = csv.reader(f)
            for row in reader:
                rows.append(row)
    return rows

def style_sheet(ws, title, headers, data_rows):
    ws.title = title
    ws.views.sheetView[0].showGridLines = True
    ws.freeze_panes = "A2"

    # Add header
    ws.append(headers)
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = THIN_BORDER
    ws.row_dimensions[1].height = 26

    # Add data rows
    for r_idx, row_data in enumerate(data_rows, start=2):
        ws.append(row_data)
        ws.row_dimensions[r_idx].height = 20
        fill_color = ZEBRA_FILL if r_idx % 2 == 0 else WHITE_FILL
        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=r_idx, column=col_idx)
            cell.font = ROW_FONT
            cell.fill = fill_color
            cell.border = THIN_BORDER
            
            # Alignments
            val_str = str(cell.value or "")
            if val_str.startswith(('JMT-', 'USR-', 'PRS-', 'TMU-', 'PSB-', 'DOA-', 'RST-')):
                cell.alignment = Alignment(horizontal="center", vertical="center")
            elif any(k in headers[col_idx-1].lower() for k in ['tanggal', 'status', 'role', 'is_', 'waktu', 'nominal']):
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")

    # Auto-fit column widths
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val = str(cell.value or '')
            if len(val) > max_len:
                max_len = len(val)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

def build_workbook():
    wb = openpyxl.Workbook()
    default_sheet = wb.active
    wb.remove(default_sheet)

    # 1. TAB JEMAAT
    jemaat_csv = read_csv("jemaat_master.csv")
    if jemaat_csv:
        jemaat_headers = jemaat_csv[0]
        jemaat_data = jemaat_csv[1:]
    else:
        jemaat_headers = ['id', 'nama_lengkap', 'nama_panggilan', 'tanggal_lahir', 'jenis_kelamin', 'kelas_kategorial', 'no_hp', 'nama_wali', 'hp_wali', 'alamat', 'status_keanggotaan', 'is_baptis', 'catatan_pastoral']
        jemaat_data = []
    ws_jemaat = wb.create_sheet()
    style_sheet(ws_jemaat, "Jemaat", jemaat_headers, jemaat_data)

    # 2. TAB PRESENSI
    presensi_headers = ['id', 'tanggal_ibadah', 'kelas_kategorial', 'jemaat_id', 'nama_jemaat', 'status_kehadiran', 'guru_pencatat', 'waktu_submit']
    sample_presensi = [
        ['PRS-0001', '2026-10-18', 'Indria (4-6 Tahun)', 'JMT-0009', 'Axel Danilla Susilo', 'Hadir', 'Cheren Masria', '2026-10-18 08:35:10'],
        ['PRS-0002', '2026-10-18', 'Indria (4-6 Tahun)', 'JMT-0010', 'Azka Pratama Tarigan', 'Hadir', 'Cheren Masria', '2026-10-18 08:35:10'],
        ['PRS-0003', '2026-10-18', 'Indria (4-6 Tahun)', 'JMT-0011', 'Benneth Evander Samuel', 'Sakit', 'Cheren Masria', '2026-10-18 08:35:10'],
        ['PRS-0004', '2026-10-18', 'Indria (4-6 Tahun)', 'JMT-0012', 'Clarissa Elvira Liu', 'Hadir', 'Cheren Masria', '2026-10-18 08:35:10'],
        ['PRS-0005', '2026-10-18', 'Indria (4-6 Tahun)', 'JMT-0013', 'Doretta Serenity Xaviera', 'Hadir', 'Cheren Masria', '2026-10-18 08:35:10'],
        ['PRS-0006', '2026-10-18', 'Gabungan Pratama (7-9 Tahun)', 'JMT-0023', 'Aleeza Vania Firmansyah', 'Hadir', 'Laudia Siwalette (Lily)', '2026-10-18 08:40:00'],
        ['PRS-0007', '2026-10-18', 'Gabungan Pratama (7-9 Tahun)', 'JMT-0024', 'Clarisa Candice Kristiawan', 'Hadir', 'Laudia Siwalette (Lily)', '2026-10-18 08:40:00'],
        ['PRS-0008', '2026-10-18', 'Madya (10-12 Tahun)', 'JMT-0035', 'Audrey Nathania Anielri', 'Hadir', 'Mexy Juliana Tambunan', '2026-10-18 08:42:00'],
        ['PRS-0009', '2026-10-18', 'Senior 2 (68+ Tahun)', 'JMT-0060', 'Bp. Thio Jan Fung', 'Hadir', 'Pdt. Hendi Tjiptamustika', '2026-10-18 08:45:00'],
        ['PRS-0010', '2026-10-18', 'Senior 2 (68+ Tahun)', 'JMT-0070', 'Ibu Tan Bok Nio', 'Alpa', 'Pdt. Hendi Tjiptamustika', '2026-10-18 08:45:00'],
    ]
    ws_presensi = wb.create_sheet()
    style_sheet(ws_presensi, "Presensi", presensi_headers, sample_presensi)

    # 3. TAB TAMU
    tamu_headers = ['id', 'tanggal_ibadah', 'kelas_kategorial', 'nama_tamu', 'no_hp', 'alamat', 'diantar_oleh', 'guru_pencatat', 'status_followup', 'catatan_pastoral']
    sample_tamu = [
        ['TMU-0001', '2026-10-18', 'Indria (4-6 Tahun)', 'Jonathan Christian', '081911223344', 'Cikole, Sukabumi', 'Doretta Serenity Xaviera', 'Cheren Masria', 'Menunggu Follow-up', 'Perlu Kunjungan Pastoral & Souvenir Sambutan Anak'],
        ['TMU-0002', '2026-10-11', 'Pemuda (20-29 Tahun)', 'Daniel Silitonga', '081399887766', 'Baros, Sukabumi', 'Raymond Owen Wijaya', 'Bp. Daniel Dean Darmawan', 'Sudah Dihubungi', 'Tertarik bergabung Komunitas Iman KOI Paulus']
    ]
    ws_tamu = wb.create_sheet()
    style_sheet(ws_tamu, "Tamu", tamu_headers, sample_tamu)

    # 4. TAB PERSEMBAHAN KELAS
    persembahan_headers = ['id', 'tanggal_ibadah', 'kelas_kategorial', 'nominal_rupiah', 'guru_pencatat', 'status_setor_bendahara', 'waktu_submit']
    sample_persembahan = [
        ['PSB-0001', '2026-10-18', 'Asuhan (0-3 Tahun)', 80000, 'Pricilia Putri Yolandi', 'Sudah Disetor', '2026-10-18 09:30:00'],
        ['PSB-0002', '2026-10-18', 'Indria (4-6 Tahun)', 185000, 'Cheren Masria Kristiana', 'Sudah Disetor', '2026-10-18 09:30:00'],
        ['PSB-0003', '2026-10-18', 'Gabungan Pratama (7-9 Tahun)', 210000, 'Laudia Siwalette (Lily)', 'Sudah Disetor', '2026-10-18 09:30:00'],
        ['PSB-0004', '2026-10-18', 'Madya (10-12 Tahun)', 165000, 'Mexy Juliana Tambunan', 'Sudah Disetor', '2026-10-18 09:30:00'],
        ['PSB-0005', '2026-10-18', 'Tunas Muda (13-15 Tahun)', 290000, 'Ibu Fung jie', 'Sudah Disetor', '2026-10-18 09:30:00'],
        ['PSB-0006', '2026-10-18', 'Remaja (16-19 Tahun)', 240000, 'Ibu Febiola Sara Ardi', 'Sudah Disetor', '2026-10-18 09:30:00'],
        ['PSB-0007', '2026-10-18', 'Pemuda (20-29 Tahun)', 310000, 'Bp. Daniel Dean Darmawan', 'Sudah Disetor', '2026-10-18 09:30:00'],
        ['PSB-0008', '2026-10-18', 'Dewasa (30-55 Tahun)', 640000, 'Pdt. Kristianto', 'Sudah Disetor', '2026-10-18 09:30:00'],
        ['PSB-0009', '2026-10-18', 'Dewasa Online (30-55 Tahun)', 200000, 'Bp. Eddy Mulyadi', 'Sudah Disetor', '2026-10-18 09:30:00'],
        ['PSB-0010', '2026-10-18', 'Senior 1 (56-67 Tahun)', 520000, 'Ibu Susan Sudarma', 'Sudah Disetor', '2026-10-18 09:30:00'],
        ['PSB-0011', '2026-10-18', 'Senior 2 (68+ Tahun)', 580000, 'Bp. Thio Jan Fung', 'Sudah Disetor', '2026-10-18 09:30:00'],
    ]
    ws_persembahan = wb.create_sheet()
    style_sheet(ws_persembahan, "Persembahan_Kelas", persembahan_headers, sample_persembahan)

    # 5. TAB POKOK DOA
    doa_headers = ['id', 'tanggal_submit', 'nama_pemohon', 'no_hp', 'kategori_doa', 'isi_doa', 'status_terjawab', 'catatan_pastoral']
    sample_doa = [
        ['DOA-0001', '2026-10-16 10:30:00', 'Bpk. S. Mulyadi', '081298114321', 'Khusus Gembala', 'Mohon doa penggembalaan untuk pemulihan pasca operasi pemasangan ring jantung di RSUD Syamsudin Sukabumi. Kiranya Tuhan memberi kekuatan dan damai sejahtera bagi keluarga.', 'Dalam Doa', 'Dijadwalkan kunjungan pastoral doa Gembala Sidang'],
        ['DOA-0002', '2026-10-15 14:00:00', 'Ibu Cheren Masria', '082210154946', 'Umum', 'Mendoakan 9 anak di Kelas Indria Sekolah Minggu agar bertumbuh dalam pengenalan akan Tuhan Yesus sejak dini, serta memohon hikmat bagi seluruh guru dalam membimbing kelas setiap Minggu.', 'Dalam Doa', 'Masuk agenda doa syafaat Rabu'],
        ['DOA-0003', '2026-10-12 18:20:00', 'Keluarga Sdri. Lily', '089657525058', 'Umum', 'Puji Tuhan doa untuk kesembuhan keponakan yang sempat demam tinggi telah dijawab Tuhan. Sekarang sudah ceria dan kembali mengikuti Sekolah Minggu.', 'Terjawab', 'Dibacakan dalam warta ucapan syukur'],
        ['DOA-0004', '2026-10-14 09:15:00', 'Anonim (Jemaat Dewasa)', '085724646100', 'Umum', 'Memohon penyertaan Tuhan bagi usaha UMKM keluarga di pasar Sukabumi di tengah pergumulan ekonomi saat ini.', 'Dalam Doa', 'Didoakan tim pendoa syafaat']
    ]
    ws_doa = wb.create_sheet()
    style_sheet(ws_doa, "Pokok_Doa", doa_headers, sample_doa)

    # 6. TAB ROSTER PELAYANAN
    roster_headers = ['id', 'tanggal_ibadah', 'sesi_ibadah', 'posisi_pelayanan', 'nama_pelayan', 'no_hp', 'status_konfirmasi', 'waktu_konfirmasi']
    sample_roster = [
        ['RST-0001', '2026-10-25', 'Ibadah Raya 1 (Pagi 07:00)', 'Pengkhotbah / Firman', 'Pdt. Kristianto, M.Th.', '081283164343', 'Bersedia', '2026-10-17 11:00:00'],
        ['RST-0002', '2026-10-25', 'Ibadah Raya 1 (Pagi 07:00)', 'Pemimpin Pujian (Liturgis)', 'Ibu Wiesye Mirnawati', '08987520893', 'Bersedia', '2026-10-17 11:30:00'],
        ['RST-0003', '2026-10-25', 'Ibadah Raya 1 (Pagi 07:00)', 'Pemusik (Keyboard)', 'Bpk. Ari', '081295859490', 'Bersedia', '2026-10-17 12:00:00'],
        ['RST-0004', '2026-10-25', 'Ibadah Raya 1 (Pagi 07:00)', 'Pemusik (Gitar Akustik)', 'Bp. Daniel Dean Darmawan', '081563199316', 'Menunggu', ''],
        ['RST-0005', '2026-10-25', 'Ibadah Raya 1 (Pagi 07:00)', 'Singer', 'Sdri. Laudia Siwalette (Lily)', '089657525058', 'Bersedia', '2026-10-17 12:15:00'],
        ['RST-0006', '2026-10-25', 'Ibadah Raya 1 (Pagi 07:00)', 'Singer', 'Sdri. Cheren Masria Kristiana', '082210154946', 'Bersedia', '2026-10-17 12:15:00'],
        ['RST-0007', '2026-10-25', 'Ibadah Raya 1 (Pagi 07:00)', 'Multimedia & Sound', 'Sdr. Reginald Alvadi', '085720710647', 'Bersedia', '2026-10-17 13:00:00'],
        ['RST-0008', '2026-10-25', 'Ibadah Raya 1 (Pagi 07:00)', 'Penyambut / Usher', 'Bp. Thio Jan Fung', '08981585969', 'Bersedia', '2026-10-17 13:30:00'],
        ['RST-0009', '2026-10-25', 'Sekolah Minggu (08:30 WIB)', 'Koordinator SM', 'Yessy Susanto', '082310626968', 'Bersedia', '2026-10-17 14:00:00'],
        ['RST-0010', '2026-10-25', 'Ibadah Raya 2 (Siang 10:00)', 'Pengkhotbah / Firman', 'Pdt. Kristianto, M.Th.', '081283164343', 'Bersedia', '2026-10-17 11:00:00']
    ]
    ws_roster = wb.create_sheet()
    style_sheet(ws_roster, "Roster_Pelayanan", roster_headers, sample_roster)

    # 7. TAB USERS
    users_csv = read_csv("users_master.csv")
    if users_csv:
        users_headers = users_csv[0]
        users_data = users_csv[1:]
    else:
        users_headers = ['id', 'no_whatsapp', 'pin_hash', 'nama', 'role', 'kelas_asuhan', 'is_active', 'last_login']
        users_data = []
    ws_users = wb.create_sheet()
    style_sheet(ws_users, "Users", users_headers, users_data)

    # Save to project
    os.makedirs(os.path.dirname(OUTPUT_PROJECT), exist_ok=True)
    wb.save(OUTPUT_PROJECT)
    print(f"SUCCESS: Saved project master spreadsheet to: {OUTPUT_PROJECT}")

    # Save to user Downloads folder
    try:
        wb.save(OUTPUT_DOWNLOADS)
        print(f"SUCCESS: Saved download copy to: {OUTPUT_DOWNLOADS}")
    except Exception as e:
        print(f"Notice: Could not save to Downloads folder directly: {e}")

if __name__ == '__main__':
    build_workbook()
