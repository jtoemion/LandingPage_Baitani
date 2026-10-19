"""
=============================================================================
GBI BAITANI SUKABUMI — MIGRATION & SEED SCRIPT
File: scripts/migrate_excel_to_sheets.py

Membaca "Presensi Juli - September 2026 per 3 bulan.xlsx" dan menghasilkan:
1. data/seeds/jemaat_master.csv (untuk tab 'Jemaat')
2. data/seeds/users_master.csv (untuk tab 'Users' dengan default PIN hash)
=============================================================================
"""

import os
import re
import csv
import hashlib
from datetime import datetime
import openpyxl

EXCEL_PATH = r"C:\Users\jtoem\Downloads\Presensi Juli - September 2026 per 3 bulan.xlsx"
OUTPUT_DIR = os.path.join(os.path.dirname(os.path.dirname(__file__)), "data", "seeds")

os.makedirs(OUTPUT_DIR, exist_ok=True)

def normalize_phone(phone_str):
    if not phone_str:
        return ""
    # If multiple numbers separated by comma or slash, take the first one
    raw_str = str(phone_str).split(',')[0].split('/')[0].strip()
    digits = re.sub(r'[^0-9+]', '', raw_str)
    if digits.startswith('+62'):
        digits = '0' + digits[3:]
    elif digits.startswith('62') and len(digits) > 9:
        digits = '0' + digits[2:]
    return digits

def parse_dob(dob_val):
    if not dob_val:
        return ""
    if isinstance(dob_val, datetime):
        return dob_val.strftime("%Y-%m-%d")
    
    val_str = str(dob_val).strip()
    # Try various formats
    formats = ["%Y-%m-%d %H:%M:%S", "%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y"]
    for fmt in formats:
        try:
            return datetime.strptime(val_str, fmt).strftime("%Y-%m-%d")
        except ValueError:
            pass
    return val_str

def format_pin_from_dob(dob_str):
    if not dob_str or len(dob_str) < 10:
        return "123456"
    try:
        dt = datetime.strptime(dob_str, "%Y-%m-%d")
        return dt.strftime("%d%m%y")
    except ValueError:
        return "123456"

def hash_pin(pin):
    return hashlib.sha256(pin.encode('utf-8')).hexdigest()

def clean_name_and_wali(name_raw, phone_raw):
    # Check if phone has (Ibu), (Bp), etc.
    wali_note = ""
    phone_str = str(phone_raw or "")
    if "(Ibu)" in phone_str or "(ibu)" in phone_str:
        wali_note = "Ibu"
    elif "(Bp)" in phone_str or "(bp)" in phone_str or "(Ayah)" in phone_str:
        wali_note = "Ayah"
    elif "(kakak)" in phone_str:
        wali_note = "Kakak"
    return str(name_raw).strip(), wali_note

def run_migration():
    if not os.path.exists(EXCEL_PATH):
        print(f"Error: Excel file not found at {EXCEL_PATH}")
        return

    wb = openpyxl.load_workbook(EXCEL_PATH, data_only=True)

    jemaat_rows = []
    user_rows = []
    seen_wa = set()
    jemaat_id_counter = 1
    user_id_counter = 1

    sheet_class_map = {
        'Asuhan': 'Asuhan (0-3 Tahun)',
        'Indria': 'Indria (4-6 Tahun)',
        'Gabungan Pratama': 'Gabungan Pratama (7-9 Tahun)',
        'Madya ': 'Madya (10-12 Tahun)',
        'Tunas Muda Putri-puta': 'Tunas Muda (13-15 Tahun)',
        'Remaja Putra-Putri': 'Remaja (16-19 Tahun)',
        'Pemuda': 'Pemuda (20-29 Tahun)',
        'Dewasa': 'Dewasa (30-55 Tahun)',
        'Dewasa on line': 'Dewasa Online (30-55 Tahun)',
        'SENIOR 1': 'Senior 1 (56-67 Tahun)',
        'SENIOR 2': 'Senior 2 (68+ Tahun)'
    }

    for sheet_name, kelas_label in sheet_class_map.items():
        if sheet_name not in wb.sheetnames:
            continue
        ws = wb[sheet_name]
        in_guru = False

        for r in range(5, ws.max_row + 1):
            row_vals = [ws.cell(row=r, column=c).value for c in range(1, 10)]
            text_line = " ".join([str(v) for v in row_vals if v is not None])

            if 'GURU' in text_line or 'PENDAMPING' in text_line:
                in_guru = True
                continue
            if 'Total' in text_line or 'PERSEMBAHAN' in text_line or 'TAMU' in text_line or 'SEKRETARIS' in text_line:
                continue

            # Identify name column (col 3 or 4)
            name_val = None
            dob_val = None
            phone_val = None

            # Check col C and D for name
            c3 = ws.cell(row=r, column=3).value
            c4 = ws.cell(row=r, column=4).value
            c5 = ws.cell(row=r, column=5).value
            c6 = ws.cell(row=r, column=6).value

            if c4 and isinstance(c4, str) and len(c4.strip()) > 2 and not c4.strip().isdigit() and 'DAFTAR' not in c4:
                name_val = c4.strip()
                dob_val = c5
                phone_val = c6
            elif c3 and isinstance(c3, str) and len(c3.strip()) > 2 and not c3.strip().isdigit() and 'DAFTAR' not in c3:
                name_val = c3.strip()
                dob_val = c4
                phone_val = c5

            if not name_val or name_val.lower() in ['nan', 'none', 'tamu', 'guru', 'murid + tamu']:
                continue

            # Skip header leftovers
            if name_val.strip().upper() in ['NAMA', 'NO', 'TGL LAHIR', 'HP/ TLP', 'JULI', 'AGUSTUS', 'SEPTEMBER', 'JUMLAH KEHADIRAN']:
                continue

            clean_name, wali_note = clean_name_and_wali(name_val, phone_val)
            clean_phone = normalize_phone(phone_val)
            dob_str = parse_dob(dob_val)

            # Determine role
            role = "Jemaat"
            if in_guru:
                role = "Guru"
                if "pdt." in clean_name.lower():
                    role = "Gembala"
                elif "sekretaris" in clean_name.lower() or "yessy" in clean_name.lower():
                    role = "Sekretariat"

            # Check if this person is child (has wali) - only pupils, never teachers
            is_child = (not in_guru) and kelas_label.startswith(('Asuhan', 'Indria', 'Gabungan Pratama', 'Madya'))
            no_hp = "" if is_child else clean_phone
            hp_wali = clean_phone if is_child else ""
            nama_wali = f"Orang Tua ({wali_note})" if is_child and wali_note else ("Orang Tua" if is_child else "")

            j_id = f"JMT-{jemaat_id_counter:04d}"
            jemaat_id_counter += 1

            jemaat_rows.append({
                'id': j_id,
                'nama_lengkap': clean_name,
                'nama_panggilan': clean_name.split()[0],
                'tanggal_lahir': dob_str,
                'jenis_kelamin': '',
                'kelas_kategorial': kelas_label,
                'no_hp': no_hp,
                'nama_wali': nama_wali,
                'hp_wali': hp_wali,
                'alamat': 'Sukabumi',
                'status_keanggotaan': 'Anggota Tetap' if not in_guru else 'Pengerja',
                'is_baptis': 'TRUE' if not is_child else 'FALSE',
                'catatan_pastoral': f"Role: {role}" if in_guru else ""
            })

            # Create User row if phone exists and not duplicate
            target_wa = no_hp or hp_wali
            if target_wa and target_wa not in seen_wa:
                seen_wa.add(target_wa)
                pin = format_pin_from_dob(dob_str)
                user_rows.append({
                    'id': f"USR-{user_id_counter:04d}",
                    'no_whatsapp': target_wa,
                    'pin_hash': hash_pin(pin),
                    'nama': clean_name if not is_child else f"{clean_name} (Keluarga)",
                    'role': role,
                    'kelas_asuhan': kelas_label if in_guru else "",
                    'is_active': 'TRUE',
                    'last_login': ''
                })
                user_id_counter += 1

    # Write Jemaat CSV
    jemaat_csv_path = os.path.join(OUTPUT_DIR, "jemaat_master.csv")
    with open(jemaat_csv_path, 'w', newline='', encoding='utf-8') as f:
        fieldnames = ['id', 'nama_lengkap', 'nama_panggilan', 'tanggal_lahir', 'jenis_kelamin', 'kelas_kategorial', 'no_hp', 'nama_wali', 'hp_wali', 'alamat', 'status_keanggotaan', 'is_baptis', 'catatan_pastoral']
        writer = csv.DictWriter(f, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(jemaat_rows)

    # Write Users CSV
    users_csv_path = os.path.join(OUTPUT_DIR, "users_master.csv")
    with open(users_csv_path, 'w', newline='', encoding='utf-8') as f:
        fieldnames = ['id', 'no_whatsapp', 'pin_hash', 'nama', 'role', 'kelas_asuhan', 'is_active', 'last_login']
        writer = csv.DictWriter(f, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(user_rows)

    print(f"SUCCESS: Migrated {len(jemaat_rows)} jemaat records -> {jemaat_csv_path}")
    print(f"SUCCESS: Generated {len(user_rows)} user login accounts -> {users_csv_path}")

if __name__ == '__main__':
    run_migration()
